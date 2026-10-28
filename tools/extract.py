"""Extract checklists from the HM SBS workbook into tools/checklists.json."""
import openpyxl,json,re,sys
wb=openpyxl.load_workbook(sys.argv[1])
roles={"Cleaner Taskoutline":("cleaner","Cleaner"),"Cashier Taskoutline":("cashier","Cashier"),"BeveragesTaskoutline ":("beverage","Beverage"),"BaristaTaskoutline":("barista","Barista"),"Manager Taskoutline ":("manager","Manager")}
clean=lambda s:re.sub(r"^\s*\d+\s*[.)]\s*","",str(s).strip())
out=[]
for ws in wb:
    rid,name=roles[ws.title]; secs=[]; cur=None; sub=None; rc=None
    for r in range(1,ws.max_row+1):
        b=ws.cell(r,2).value
        if b is None: continue
        hc=next((c for c in range(10,ws.max_column+1) if "Received" in str(ws.cell(r,c).value or "")),None)
        if hc:
            rc=hc;cur={"name":str(b).strip(),"tasks":[]};secs.append(cur);sub=None;continue
        if cur is None or rc is None: continue
        if ws.cell(r,rc).value is not None:
            t=clean(b); cur["tasks"].append(f"{sub}: {t}" if sub else t)
        else:
            s=str(b).strip()
            if len(s)<40 and not s.lower().startswith(("note","it is","serve")): sub=s
    secs=[s for s in secs if s["tasks"]]
    out.append({"id":rid,"name":name,"sections":secs})
    print(name,sum(len(s["tasks"]) for s in secs),[(s["name"][:25],len(s["tasks"])) for s in secs])
json.dump(out,open("tools/checklists.json","w"),indent=1,ensure_ascii=False)
